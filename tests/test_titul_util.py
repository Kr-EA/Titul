import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from docx import Document
from PyQt6.QtWidgets import QApplication, QMessageBox
from titul_util import MainWindow, SettingsDialog, TitulGenerator


class TitulTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.config_path = self.directory / "settings" / "config.json"
        self.generator = TitulGenerator(config_path=self.config_path)

    def test_settings_save_and_reload(self):
        dialog = SettingsDialog(self.generator)
        dialog.fields["me"].setText("Тестов Т.Т.")
        dialog.fields["my_caf"].setCurrentText("МК1")
        dialog.fields["group_num"].setValue(3)
        dialog.fields["degree"].setCurrentText("Магистр")
        dialog.save()
        self.assertEqual(dialog.result(), dialog.DialogCode.Accepted)
        loaded = TitulGenerator(config_path=self.config_path)
        self.assertEqual(loaded.defaults["author"], "Тестов Т.Т.")
        self.assertEqual(loaded.defaults["caf"], "МК1")
        self.assertTrue(loaded.defaults["group"].endswith("3М"))
        self.assertEqual(json.loads(self.config_path.read_text(encoding="utf-8"))["me"], "Тестов Т.Т.")

    def test_cancel_leaves_config_unchanged(self):
        original = self.config_path.read_bytes()
        dialog = SettingsDialog(self.generator)
        dialog.fields["me"].setText("Не сохранять")
        dialog.reject()
        self.assertEqual(self.config_path.read_bytes(), original)

    def test_invalid_settings_and_write_failure_preserve_config(self):
        original = self.config_path.read_bytes()
        config = dict(self.generator.config, me="")
        with self.assertRaises(ValueError):
            self.generator.save_config(config)
        config["me"] = "Новое имя"
        with patch("titul_util.os.replace", side_effect=PermissionError("Нет доступа")):
            with self.assertRaises(PermissionError):
                self.generator.save_config(config)
        self.assertEqual(self.config_path.read_bytes(), original)
        self.assertNotEqual(self.generator.defaults["author"], "Новое имя")
        self.assertEqual(list(self.config_path.parent.glob("*.tmp")), [])

    def test_form_creates_document_using_saved_settings(self):
        self.generator.save_config(dict(self.generator.config, me="Тестов Т.Т."))
        window = MainWindow(self.generator, self.directory)
        self.addCleanup(window.close)
        for key, value in {"w_title": "ЛР №1", "subject": "Операционные системы", "tutor": "АЯГ"}.items():
            window.inputs[key].setText(value)
        with patch.object(QMessageBox, "information") as success, patch.object(QMessageBox, "critical") as error:
            window.create_document()
            self.assertFalse(error.called, error.call_args)
            self.assertTrue(success.called)
        files = list(self.directory.glob("*.docx"))
        self.assertEqual(len(files), 1)
        document = Document(files[0])
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
        self.assertIn("Тестов Т.Т.", text)
        self.assertIn("ЛАБОРАТОРНАЯ РАБОТА №1", text)
        self.assertNotIn("{{", text)

    def test_installed_app_uses_user_config(self):
        user_dir = self.directory / "user-config"
        with patch("titul_util.sys.frozen", True, create=True), patch(
            "titul_util.QStandardPaths.writableLocation", return_value=str(user_dir)
        ):
            installed = TitulGenerator()
            self.assertEqual(installed.config_path, user_dir / "config.json")
            installed.save_config(dict(installed.config, me="Пользователь П.П."))
            restarted = TitulGenerator()
            self.assertEqual(restarted.defaults["author"], "Пользователь П.П.")

    def test_user_config_not_overwritten_on_startup(self):
        self.generator.save_config(dict(self.generator.config, me="Сохранённый С.С."))
        TitulGenerator(config_path=self.config_path)
        self.assertEqual(json.loads(self.config_path.read_text(encoding="utf-8"))["me"], "Сохранённый С.С.")


if __name__ == "__main__":
    unittest.main()
